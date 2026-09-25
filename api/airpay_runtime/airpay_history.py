"""Фильтры и выгрузка сохранённой истории Airpay без транспорта поставщика."""
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field, ValidationError, model_validator
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill

from airpay_runtime.airpay_contract import transaction_outcome

MSK = ZoneInfo('Europe/Moscow')
EXPORT_LIMIT = 10000
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
STATES = {'prepared': 'Подготовка', 'checking': 'Проверяется', 'check_pending': 'Ожидает проверки',
          'check_failed': 'Проверка не пройдена', 'checked': 'Проверен', 'processing': 'Исход оплаты неизвестен',
          'paid': 'Оплачен', 'failed': 'Отказ в оплате'}


class HistoryFilters(BaseModel):
    q: str = Field(default='', max_length=200)
    status: Literal['', 'prepared', 'checking', 'check_pending', 'check_failed', 'checked', 'processing',
                    'paid', 'failed', 'completed', 'awaiting_voucher', 'requires_attention'] = ''
    kind: Literal['', 'voucher', 'topup'] = ''
    date_from: date | None = None
    date_to: date | None = None

    @model_validator(mode='after')
    def check_dates(self):
        # Включаем весь последний день, не принимая перевёрнутый или переполняющий дату диапазон.
        if self.date_to == date.max or (self.date_from and self.date_to and self.date_from > self.date_to):
            raise ValueError('Некорректный диапазон дат')
        self.q = self.q.strip()
        return self


def read_history_filters(request: Request):
    # Один разбор параметров используется прямым API, Hub и архивом CRM.
    try:
        return HistoryFilters.model_validate({key: request.query_params[key] for key in HistoryFilters.model_fields if request.query_params.get(key)})
    except ValidationError:
        raise HTTPException(422, 'Некорректные фильтры истории Airpay: проверьте даты, состояние и поиск (до 200 символов)') from None


def history_where(owner, filters):
    # Имена колонок заданы кодом; пользовательские значения передаются только параметрами SQL.
    clauses, params = ['created_by=%s'], [owner]
    if filters.q:
        escaped = filters.q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
        columns = ('agent_transaction_id::text', 'provider_transaction_id', 'service_id', 'service_title', 'account')
        clauses.append('(' + ' OR '.join(f"{column} ILIKE %s ESCAPE E'\\\\'" for column in columns) + ')')
        params.extend(['%' + escaped + '%'] * len(columns))
    if filters.kind:
        clauses.append('purchase_kind=%s')
        params.append(filters.kind)
    for value, operator in ((filters.date_from, '>='), (filters.date_to, '<')):
        if value:
            clauses.append('created_at ' + operator + ' %s')
            params.append(datetime.combine(value + timedelta(days=1) if operator == '<' else value, time.min, MSK))
    code = "(pin_ciphertext IS NOT NULL OR COALESCE(pin_code, '') <> '')"
    if filters.status == 'completed':
        clauses.append(f"(state='paid' AND (purchase_kind='topup' OR {code}))")
    elif filters.status == 'awaiting_voucher':
        clauses.append(f"(state='paid' AND purchase_kind='voucher' AND NOT {code})")
    elif filters.status == 'requires_attention':
        clauses.append('requires_attention=true')
    elif filters.status:
        clauses.append('state=%s')
        params.append(filters.status)
    return ' AND '.join(clauses), params


def export_history(repository, owner, filters):
    # Одна выборка экспортирует весь результат; превышение лимита не скрывается усечённым файлом.
    rows = repository.history(owner, EXPORT_LIMIT + 1, 0, filters)
    if len(rows) > EXPORT_LIMIT:
        raise HTTPException(422, 'Найдено больше 10 000 операций. Уточните фильтры для выгрузки.')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Airpay'
    sheet.append(['Номер операции', 'Создана (МСК)', 'Услуга', 'Код услуги', 'Аккаунт', 'Вид покупки',
                  'Состояние', 'Закупочная сумма', 'Валюта', 'Номер у поставщика', 'Результат',
                  'Требует разбора', 'Причина разбора', 'Группа покупки', 'Позиция в группе', 'Заменена подготовкой'])
    for row in rows:
        outcome = transaction_outcome(row)
        batch = (row.get('service_snapshot') or {}).get('_airpay_batch') or {}
        created = row.get('created_at')
        if isinstance(created, datetime):
            created = created.astimezone(MSK).strftime('%Y-%m-%d %H:%M:%S')
        values = [row['agent_transaction_id'], created, row.get('service_title'), row.get('service_id'), row.get('account'),
                  'Ваучер' if row['purchase_kind'] == 'voucher' else 'Пополнение', STATES.get(row['state'], row['state']),
                  row.get('amount'), row.get('currency'), row.get('provider_transaction_id'),
                  {'ready': 'Код сохранён', 'pending': 'Ожидает кода', 'unavailable': 'Не получен', 'not_required': 'Код не требуется'}[outcome['result_state']],
                  'Да' if outcome['requires_attention'] else 'Нет', outcome['attention_reason'], batch.get('key'), (batch['index'] + 1) if 'index' in batch else '',
                  'Да' if row.get('replacement_key') else 'Нет']
        sheet.append([''] * len(values))
        for cell, value in zip(sheet[sheet.max_row], values):
            # Текст сохраняет длинные ID и точные суммы, а значения с «=» не становятся формулами.
            cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(value) if value is not None else '')[:32767]
            cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='273248')
        sheet.column_dimensions[cell.column_letter].width = 25 if cell.column != 3 else 45
    buffer = BytesIO()
    workbook.save(buffer)
    return Response(buffer.getvalue(), media_type=XLSX_TYPE,
                    headers={'Content-Disposition': 'attachment; filename="airpay-history.xlsx"', 'Cache-Control': 'no-store'})
